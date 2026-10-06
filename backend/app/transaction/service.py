"""Transaction service - redesigned for ClickHouse from first principles."""

import csv
import logging
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from typing import Literal, Optional

from openpyxl import Workbook
from openpyxl.styles import Font

from app.clickhouse.client import query
from app.clickhouse.queries import _project_filter
from app.transaction.models import (
    RentalContractListResponse,
    RentalContractResponse,
    TransactionListResponse,
    TransactionResponse,
)

logger = logging.getLogger(__name__)

MAX_TRANSACTION_EXPORT_ROWS = 50_000
SPREADSHEET_FORMULA_PREFIXES = ("=", "+", "-", "@")
SPREADSHEET_CONTROL_PREFIXES = ("\t", "\r", "\n")


def _date_only(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return value


def _active_export_filters(filters: dict[str, object]) -> str:
    active_filters = [
        f"{key}={value}"
        for key, value in filters.items()
        if value is not None and value != "" and value != "All"
    ]
    return ", ".join(active_filters) if active_filters else "None"


def _add_export_metadata(
    workbook: Workbook,
    *,
    title: str,
    generated_at: datetime,
    row_count: int,
    filters: dict[str, object],
    source: str,
) -> None:
    metadata = workbook.create_sheet("Metadata", 0)
    rows = [
        ("Publisher", "Dubai Land Department"),
        ("Dataset", title),
        ("Generated at", generated_at.isoformat()),
        ("Data source", source),
        ("Rows exported", row_count),
        ("Maximum filtered rows", MAX_TRANSACTION_EXPORT_ROWS),
        ("Filters applied", _active_export_filters(filters)),
    ]
    for row in rows:
        metadata.append(row)
    for cell in metadata["A"]:
        cell.font = Font(bold=True)


def _add_export_header(
    worksheet,
    *,
    title: str,
    generated_at: datetime,
    source: str,
) -> None:
    worksheet.append(["Dubai Land Department"])
    worksheet.append([title])
    worksheet.append([f"Generated at {generated_at.isoformat()}"])
    worksheet.append([f"Source: {source}"])
    worksheet.append([])
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A2"].font = Font(bold=True)


def _sanitize_spreadsheet_value(value: object) -> object:
    if not isinstance(value, str) or not value:
        return value

    stripped = value.lstrip()
    if value.startswith(SPREADSHEET_CONTROL_PREFIXES) or (
        stripped and stripped.startswith(SPREADSHEET_FORMULA_PREFIXES)
    ):
        return f"'{value}"
    return value


def _sanitize_spreadsheet_row(row: list[object]) -> list[object]:
    return [_sanitize_spreadsheet_value(value) for value in row]


def _rows_to_csv(headers: list[str], rows: list[list[object]]) -> BytesIO:
    csv_text = StringIO()
    writer = csv.writer(csv_text)
    writer.writerow(headers)
    writer.writerows(_sanitize_spreadsheet_row(row) for row in rows)

    output = BytesIO()
    output.write(csv_text.getvalue().encode("utf-8-sig"))
    output.seek(0)
    return output


async def fetch_sales_transactions(
    limit: int,
    offset: int,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    rooms: Optional[str] = None,
    registration_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    building: Optional[str] = None,
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    price_min: Optional[float] = None,
    price_max: Optional[float] = None,
    price_per_sqm_min: Optional[float] = None,
    price_per_sqm_max: Optional[float] = None,
) -> TransactionListResponse:
    """
    Fetch paginated sales transactions from ClickHouse.

    Optimized query using sort key (property_usage_en, instance_date, transaction_id).
    """
    # Build WHERE clause — always enforce Sales only (separate endpoint exists for mortgages/gifts)
    where_clauses = ["trans_group_en = 'Sales'"]
    params: dict = {}

    if property_usage:
        where_clauses.append("property_usage_en = {property_usage:String}")
        params["property_usage"] = property_usage

    if property_type:
        where_clauses.append("property_type_en = {property_type:String}")
        params["property_type"] = property_type

    if property_sub_type:
        where_clauses.append("property_sub_type_en = {property_sub_type:String}")
        params["property_sub_type"] = property_sub_type

    if rooms:
        where_clauses.append("rooms_en = {rooms:String}")
        params["rooms"] = rooms

    if registration_type:
        where_clauses.append("reg_type_en = {registration_type:String}")
        params["registration_type"] = registration_type

    if area:
        where_clauses.append("area_name_en = {area:String}")
        params["area"] = area

    if project:
        where_clauses.append(_project_filter(filter_type))
        params["project_name"] = project + " - %" if filter_type == "virtual_master" else project

    if building:
        where_clauses.append("building_name_en = {building:String}")
        params["building"] = building

    if timeframe_days is not None:
        where_clauses.append("instance_date >= today() - {timeframe_days:Int32}")
        params["timeframe_days"] = timeframe_days

    if area_min is not None:
        where_clauses.append("procedure_area >= {area_min:Decimal64(2)}")
        params["area_min"] = area_min

    if area_max is not None:
        where_clauses.append("procedure_area <= {area_max:Decimal64(2)}")
        params["area_max"] = area_max

    if price_min is not None:
        where_clauses.append("actual_worth >= {price_min:Decimal64(2)}")
        params["price_min"] = price_min

    if price_max is not None:
        where_clauses.append("actual_worth <= {price_max:Decimal64(2)}")
        params["price_max"] = price_max

    if price_per_sqm_min is not None:
        where_clauses.append("meter_sale_price >= {price_per_sqm_min:Decimal64(2)}")
        params["price_per_sqm_min"] = price_per_sqm_min

    if price_per_sqm_max is not None:
        where_clauses.append("meter_sale_price <= {price_per_sqm_max:Decimal64(2)}")
        params["price_per_sqm_max"] = price_per_sqm_max

    where_clause = " AND ".join(where_clauses)

    # Count total matching transactions
    count_sql = f"""
        SELECT count() as total
        FROM ch_transactions
        WHERE {where_clause}
    """
    count_result = await query(count_sql, params)
    total = count_result[0]["total"] if count_result else 0

    # Fetch transactions
    params["limit"] = limit
    params["offset"] = offset

    transactions_sql = f"""
        SELECT
            transaction_id,
            instance_date,
            property_usage_en,
            trans_group_en,
            procedure_name_en,
            property_type_en,
            property_sub_type_en,
            rooms_en,
            procedure_area,
            actual_worth,
            meter_sale_price,
            reg_type_en,
            area_name_en,
            project_name_en,
            building_name_en,
            nearest_landmark_en,
            nearest_metro_en,
            nearest_mall_en
        FROM ch_transactions
        WHERE {where_clause}
        ORDER BY instance_date DESC, transaction_id
        LIMIT {{limit:Int32}}
        OFFSET {{offset:Int32}}
    """
    transactions_result = await query(transactions_sql, params)

    # Convert to response models
    items = [
        TransactionResponse(
            transaction_id=row["transaction_id"],
            instance_date=_date_only(row["instance_date"]),
            trans_group_en=row["trans_group_en"],
            procedure_name_en=row["procedure_name_en"],
            property_type_en=row["property_type_en"],
            property_sub_type_en=row["property_sub_type_en"],
            rooms_en=row["rooms_en"],
            area_name_en=row["area_name_en"],
            project_name_en=row["project_name_en"],
            building_name_en=row["building_name_en"],
            procedure_area=float(row["procedure_area"]) if row["procedure_area"] else None,
            actual_worth=float(row["actual_worth"]) if row["actual_worth"] else None,
            meter_sale_price=float(row["meter_sale_price"]) if row["meter_sale_price"] else None,
            reg_type_en=row["reg_type_en"],
            nearest_landmark_en=row["nearest_landmark_en"],
            nearest_metro_en=row["nearest_metro_en"],
            nearest_mall_en=row["nearest_mall_en"],
        )
        for row in transactions_result
    ]

    return TransactionListResponse(transactions=items, total=total, limit=limit, offset=offset)


async def fetch_rental_contracts(
    limit: int,
    offset: int,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    tenant_type: Optional[str] = None,
    contract_reg_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    rent_min: Optional[float] = None,
    rent_max: Optional[float] = None,
) -> RentalContractListResponse:
    """
    Fetch paginated rental contracts from ClickHouse.

    Optimized query using sort key (property_usage_en, contract_start_date, contract_id).
    """
    # Build WHERE clause
    where_clauses = []
    params = {}

    if property_usage:
        where_clauses.append("property_usage_en = {property_usage:String}")
        params["property_usage"] = property_usage

    if property_type:
        where_clauses.append("ejari_property_type_en = {property_type:String}")
        params["property_type"] = property_type

    if property_sub_type:
        where_clauses.append("ejari_property_sub_type_en = {property_sub_type:String}")
        params["property_sub_type"] = property_sub_type

    if tenant_type:
        where_clauses.append("tenant_type_en = {tenant_type:String}")
        params["tenant_type"] = tenant_type

    if contract_reg_type:
        where_clauses.append("contract_reg_type_en = {contract_reg_type:String}")
        params["contract_reg_type"] = contract_reg_type

    if area:
        where_clauses.append("area_name_en = {area:String}")
        params["area"] = area

    if project:
        where_clauses.append(_project_filter(filter_type))
        params["project_name"] = project + " - %" if filter_type == "virtual_master" else project

    if timeframe_days is not None:
        where_clauses.append("contract_start_date >= today() - {timeframe_days:Int32}")
        params["timeframe_days"] = timeframe_days

    if area_min is not None:
        where_clauses.append("actual_area >= {area_min:Decimal64(2)}")
        params["area_min"] = area_min

    if area_max is not None:
        where_clauses.append("actual_area <= {area_max:Decimal64(2)}")
        params["area_max"] = area_max

    if rent_min is not None:
        where_clauses.append("annual_amount >= {rent_min:Decimal64(2)}")
        params["rent_min"] = rent_min

    if rent_max is not None:
        where_clauses.append("annual_amount <= {rent_max:Decimal64(2)}")
        params["rent_max"] = rent_max

    where_clause = " AND ".join(where_clauses) if where_clauses else "1=1"

    # Count total matching contracts
    count_sql = f"""
        SELECT count() as total
        FROM ch_rent_contracts
        WHERE {where_clause}
    """
    count_result = await query(count_sql, params)
    total = count_result[0]["total"] if count_result else 0

    # Fetch rental contracts
    params["limit"] = limit
    params["offset"] = offset

    contracts_sql = f"""
        SELECT
            contract_id,
            contract_start_date,
            contract_end_date,
            property_usage_en,
            contract_reg_type_en,
            ejari_property_type_en,
            ejari_property_sub_type_en,
            tenant_type_en,
            actual_area,
            annual_amount,
            contract_amount,
            area_name_en,
            project_name_en,
            nearest_landmark_en,
            nearest_metro_en,
            nearest_mall_en,
            no_of_prop,
            is_free_hold
        FROM ch_rent_contracts
        WHERE {where_clause}
        ORDER BY contract_start_date DESC, contract_id
        LIMIT {{limit:Int32}}
        OFFSET {{offset:Int32}}
    """
    contracts_result = await query(contracts_sql, params)

    # Convert to response models
    items = [
        RentalContractResponse(
            contract_id=row["contract_id"],
            contract_start_date=row["contract_start_date"],
            contract_end_date=row["contract_end_date"],
            ejari_property_type_en=row["ejari_property_type_en"],
            ejari_property_sub_type_en=row["ejari_property_sub_type_en"],
            tenant_type_en=row["tenant_type_en"],
            contract_reg_type_en=row["contract_reg_type_en"],
            area_name_en=row["area_name_en"],
            project_name_en=row["project_name_en"],
            actual_area=float(row["actual_area"]) if row["actual_area"] else None,
            annual_amount=float(row["annual_amount"]) if row["annual_amount"] else None,
            contract_amount=float(row["contract_amount"]) if row["contract_amount"] else None,
            no_of_prop=row["no_of_prop"],
            is_free_hold=bool(row["is_free_hold"]),
            nearest_landmark_en=row["nearest_landmark_en"],
            nearest_metro_en=row["nearest_metro_en"],
            nearest_mall_en=row["nearest_mall_en"],
        )
        for row in contracts_result
    ]

    return RentalContractListResponse(contracts=items, total=total, limit=limit, offset=offset)


async def export_sales_transactions(
    transaction_ids: Optional[str] = None,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    rooms: Optional[str] = None,
    registration_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    building: Optional[str] = None,
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    price_min: Optional[float] = None,
    price_max: Optional[float] = None,
    price_per_sqm_min: Optional[float] = None,
    price_per_sqm_max: Optional[float] = None,
    file_format: Literal["xlsx", "csv"] = "xlsx",
) -> BytesIO:
    """
    Export sales transactions to Excel or CSV.

    Limit: 50,000 rows max (or specific transaction_ids if provided).
    """
    params = {}
    filters = {
        "transaction_ids": transaction_ids,
        "property_usage": property_usage,
        "property_type": property_type,
        "property_sub_type": property_sub_type,
        "rooms": rooms,
        "registration_type": registration_type,
        "area": area,
        "project": project,
        "filter_type": filter_type,
        "building": building,
        "timeframe_days": timeframe_days,
        "area_min": area_min,
        "area_max": area_max,
        "price_min": price_min,
        "price_max": price_max,
        "price_per_sqm_min": price_per_sqm_min,
        "price_per_sqm_max": price_per_sqm_max,
    }

    # If specific IDs provided, use them (no limit)
    if transaction_ids:
        ids_list = [tid.strip() for tid in transaction_ids.split(",")]
        # Build IN clause with parameter binding
        placeholders = ", ".join([f"{{id_{i}:String}}" for i in range(len(ids_list))])
        where_clause = f"transaction_id IN ({placeholders})"
        for i, tid in enumerate(ids_list):
            params[f"id_{i}"] = tid
    else:
        # Build WHERE clause from filters — always enforce Sales only
        where_clauses = ["trans_group_en = 'Sales'"]

        if property_usage:
            where_clauses.append("property_usage_en = {property_usage:String}")
            params["property_usage"] = property_usage

        if property_type:
            where_clauses.append("property_type_en = {property_type:String}")
            params["property_type"] = property_type

        if property_sub_type:
            where_clauses.append("property_sub_type_en = {property_sub_type:String}")
            params["property_sub_type"] = property_sub_type

        if rooms:
            where_clauses.append("rooms_en = {rooms:String}")
            params["rooms"] = rooms

        if registration_type:
            where_clauses.append("reg_type_en = {registration_type:String}")
            params["registration_type"] = registration_type

        if area:
            where_clauses.append("area_name_en = {area:String}")
            params["area"] = area

        if project:
            where_clauses.append(_project_filter(filter_type))
            params["project_name"] = (
                project + " - %" if filter_type == "virtual_master" else project
            )

        if building:
            where_clauses.append("building_name_en = {building:String}")
            params["building"] = building

        if timeframe_days is not None:
            where_clauses.append("instance_date >= today() - {timeframe_days:Int32}")
            params["timeframe_days"] = timeframe_days

        if area_min is not None:
            where_clauses.append("procedure_area >= {area_min:Decimal64(2)}")
            params["area_min"] = area_min

        if area_max is not None:
            where_clauses.append("procedure_area <= {area_max:Decimal64(2)}")
            params["area_max"] = area_max

        if price_min is not None:
            where_clauses.append("actual_worth >= {price_min:Decimal64(2)}")
            params["price_min"] = price_min

        if price_max is not None:
            where_clauses.append("actual_worth <= {price_max:Decimal64(2)}")
            params["price_max"] = price_max

        if price_per_sqm_min is not None:
            where_clauses.append("meter_sale_price >= {price_per_sqm_min:Decimal64(2)}")
            params["price_per_sqm_min"] = price_per_sqm_min

        if price_per_sqm_max is not None:
            where_clauses.append("meter_sale_price <= {price_per_sqm_max:Decimal64(2)}")
            params["price_per_sqm_max"] = price_per_sqm_max

        where_clause = " AND ".join(where_clauses)

    # Fetch transactions (limit 50,000 if using filters)
    if transaction_ids:
        limit_clause = ""
    else:
        params["export_limit"] = MAX_TRANSACTION_EXPORT_ROWS
        limit_clause = "LIMIT {export_limit:Int32}"

    export_sql = f"""
        SELECT
            transaction_id,
            instance_date,
            property_usage_en,
            trans_group_en,
            procedure_name_en,
            property_type_en,
            property_sub_type_en,
            rooms_en,
            procedure_area,
            actual_worth,
            meter_sale_price,
            reg_type_en,
            area_name_en,
            project_name_en,
            building_name_en,
            nearest_landmark_en,
            nearest_metro_en,
            nearest_mall_en
        FROM ch_transactions
        WHERE {where_clause}
        ORDER BY instance_date DESC, transaction_id
        {limit_clause}
    """
    results = await query(export_sql, params)

    headers = [
        "Transaction ID",
        "Date",
        "Usage",
        "Trans Group",
        "Procedure",
        "Type",
        "Sub Type",
        "Rooms",
        "Area (sqm)",
        "Price",
        "Price/sqm",
        "Registration Type",
        "Area",
        "Project",
        "Building",
        "Landmark",
        "Metro",
        "Mall",
    ]
    rows = [
        [
            row["transaction_id"],
            row["instance_date"],
            row["property_usage_en"],
            row["trans_group_en"],
            row["procedure_name_en"],
            row["property_type_en"],
            row["property_sub_type_en"],
            row["rooms_en"],
            float(row["procedure_area"]) if row["procedure_area"] else None,
            float(row["actual_worth"]) if row["actual_worth"] else None,
            float(row["meter_sale_price"]) if row["meter_sale_price"] else None,
            row["reg_type_en"],
            row["area_name_en"],
            row["project_name_en"],
            row["building_name_en"],
            row["nearest_landmark_en"],
            row["nearest_metro_en"],
            row["nearest_mall_en"],
        ]
        for row in results
    ]

    if file_format == "csv":
        return _rows_to_csv(headers, rows)

    generated_at = datetime.now(UTC)
    source = "DLD transaction registry"

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Transactions"
    _add_export_metadata(
        wb,
        title="Sales Transactions",
        generated_at=generated_at,
        row_count=len(rows),
        filters=filters,
        source=source,
    )
    _add_export_header(ws, title="Sales Transactions", generated_at=generated_at, source=source)
    ws.append(headers)

    # Style headers
    for cell in ws[6]:
        cell.font = Font(bold=True)

    # Add data rows
    for row in rows:
        ws.append(_sanitize_spreadsheet_row(row))

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_rental_contracts(
    contract_ids: Optional[str] = None,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    tenant_type: Optional[str] = None,
    contract_reg_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    rent_min: Optional[float] = None,
    rent_max: Optional[float] = None,
    file_format: Literal["xlsx", "csv"] = "xlsx",
) -> BytesIO:
    """
    Export rental contracts to Excel or CSV.

    Limit: 50,000 rows max (or specific contract_ids if provided).
    """
    params = {}
    filters = {
        "contract_ids": contract_ids,
        "property_usage": property_usage,
        "property_type": property_type,
        "property_sub_type": property_sub_type,
        "tenant_type": tenant_type,
        "contract_reg_type": contract_reg_type,
        "area": area,
        "project": project,
        "filter_type": filter_type,
        "timeframe_days": timeframe_days,
        "area_min": area_min,
        "area_max": area_max,
        "rent_min": rent_min,
        "rent_max": rent_max,
    }

    # If specific IDs provided, use them (no limit)
    if contract_ids:
        ids_list = [cid.strip() for cid in contract_ids.split(",")]
        # Build IN clause with parameter binding
        placeholders = ", ".join([f"{{id_{i}:String}}" for i in range(len(ids_list))])
        where_clause = f"contract_id IN ({placeholders})"
        for i, cid in enumerate(ids_list):
            params[f"id_{i}"] = cid
    else:
        # Build WHERE clause from filters
        where_clauses = []

        if property_usage:
            where_clauses.append("property_usage_en = {property_usage:String}")
            params["property_usage"] = property_usage

        if property_type:
            where_clauses.append("ejari_property_type_en = {property_type:String}")
            params["property_type"] = property_type

        if property_sub_type:
            where_clauses.append("ejari_property_sub_type_en = {property_sub_type:String}")
            params["property_sub_type"] = property_sub_type

        if tenant_type:
            where_clauses.append("tenant_type_en = {tenant_type:String}")
            params["tenant_type"] = tenant_type

        if contract_reg_type:
            where_clauses.append("contract_reg_type_en = {contract_reg_type:String}")
            params["contract_reg_type"] = contract_reg_type

        if area:
            where_clauses.append("area_name_en = {area:String}")
            params["area"] = area

        if project:
            where_clauses.append(_project_filter(filter_type))
            params["project_name"] = (
                project + " - %" if filter_type == "virtual_master" else project
            )

        if timeframe_days is not None:
            where_clauses.append("contract_start_date >= today() - {timeframe_days:Int32}")
            params["timeframe_days"] = timeframe_days

        if area_min is not None:
            where_clauses.append("actual_area >= {area_min:Decimal64(2)}")
            params["area_min"] = area_min

        if area_max is not None:
            where_clauses.append("actual_area <= {area_max:Decimal64(2)}")
            params["area_max"] = area_max

        if rent_min is not None:
            where_clauses.append("annual_amount >= {rent_min:Decimal64(2)}")
            params["rent_min"] = rent_min

        if rent_max is not None:
            where_clauses.append("annual_amount <= {rent_max:Decimal64(2)}")
            params["rent_max"] = rent_max

        where_clause = " AND ".join(where_clauses) if where_clauses else "1=1"

    # Fetch contracts (limit 50,000 if using filters)
    if contract_ids:
        limit_clause = ""
    else:
        params["export_limit"] = MAX_TRANSACTION_EXPORT_ROWS
        limit_clause = "LIMIT {export_limit:Int32}"

    export_sql = f"""
        SELECT
            contract_id,
            contract_start_date,
            contract_end_date,
            property_usage_en,
            contract_reg_type_en,
            ejari_property_type_en,
            ejari_property_sub_type_en,
            tenant_type_en,
            actual_area,
            annual_amount,
            contract_amount,
            area_name_en,
            project_name_en,
            nearest_landmark_en,
            nearest_metro_en,
            nearest_mall_en,
            no_of_prop,
            is_free_hold
        FROM ch_rent_contracts
        WHERE {where_clause}
        ORDER BY contract_start_date DESC, contract_id
        {limit_clause}
    """
    results = await query(export_sql, params)

    headers = [
        "Contract ID",
        "Start Date",
        "End Date",
        "Usage",
        "Reg Type",
        "Type",
        "Sub Type",
        "Tenant Type",
        "Area (sqm)",
        "Annual Rent",
        "Contract Amount",
        "Area",
        "Project",
        "Landmark",
        "Metro",
        "Mall",
        "No. Properties",
        "Freehold",
    ]
    rows = [
        [
            row["contract_id"],
            row["contract_start_date"],
            row["contract_end_date"],
            row["property_usage_en"],
            row["contract_reg_type_en"],
            row["ejari_property_type_en"],
            row["ejari_property_sub_type_en"],
            row["tenant_type_en"],
            float(row["actual_area"]) if row["actual_area"] else None,
            float(row["annual_amount"]) if row["annual_amount"] else None,
            float(row["contract_amount"]) if row["contract_amount"] else None,
            row["area_name_en"],
            row["project_name_en"],
            row["nearest_landmark_en"],
            row["nearest_metro_en"],
            row["nearest_mall_en"],
            row["no_of_prop"],
            "Yes" if row["is_free_hold"] else "No",
        ]
        for row in results
    ]

    if file_format == "csv":
        return _rows_to_csv(headers, rows)

    generated_at = datetime.now(UTC)
    source = "DLD Ejari rental contracts"

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Rental Contracts"
    _add_export_metadata(
        wb,
        title="Rental Contracts",
        generated_at=generated_at,
        row_count=len(rows),
        filters=filters,
        source=source,
    )
    _add_export_header(ws, title="Rental Contracts", generated_at=generated_at, source=source)
    ws.append(headers)

    # Style headers
    for cell in ws[6]:
        cell.font = Font(bold=True)

    # Add data rows
    for row in rows:
        ws.append(_sanitize_spreadsheet_row(row))

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
