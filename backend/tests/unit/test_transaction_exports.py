from openpyxl import load_workbook

from app.transaction import service as transaction_service

SALES_ROW = {
    "transaction_id": "T-1",
    "instance_date": "2026-05-01",
    "property_usage_en": "Residential",
    "trans_group_en": "Sales",
    "procedure_name_en": "Sell",
    "property_type_en": "Unit",
    "property_sub_type_en": "Flat",
    "rooms_en": "2 B/R",
    "procedure_area": 100,
    "actual_worth": 1_500_000,
    "meter_sale_price": 15_000,
    "reg_type_en": "Ready",
    "area_name_en": "Dubai Marina",
    "project_name_en": "Marina Tower",
    "building_name_en": "Marina Tower A",
    "nearest_landmark_en": "Marina Mall",
    "nearest_metro_en": "DMCC",
    "nearest_mall_en": "Dubai Marina Mall",
}

RENTAL_ROW = {
    "contract_id": "R-1",
    "contract_start_date": "2026-05-01",
    "contract_end_date": "2027-04-30",
    "property_usage_en": "Residential",
    "contract_reg_type_en": "New",
    "ejari_property_type_en": "Unit",
    "ejari_property_sub_type_en": "Flat",
    "tenant_type_en": "Person",
    "actual_area": 100,
    "annual_amount": 120_000,
    "contract_amount": 120_000,
    "area_name_en": "Dubai Marina",
    "project_name_en": "Marina Tower",
    "nearest_landmark_en": "Marina Mall",
    "nearest_metro_en": "DMCC",
    "nearest_mall_en": "Dubai Marina Mall",
    "no_of_prop": 1,
    "is_free_hold": True,
}


async def test_sales_excel_export_has_dld_metadata_and_50000_row_limit(monkeypatch):
    seen = {}

    async def fake_query(sql, params):
        seen["sql"] = sql
        seen["params"] = params
        return [SALES_ROW]

    monkeypatch.setattr(transaction_service, "query", fake_query)

    stream = await transaction_service.export_sales_transactions(area="Dubai Marina")

    assert "LIMIT {export_limit:Int32}" in seen["sql"]
    assert seen["params"]["export_limit"] == transaction_service.MAX_TRANSACTION_EXPORT_ROWS

    workbook = load_workbook(stream)
    assert workbook.sheetnames[:2] == ["Metadata", "Sales Transactions"]
    assert workbook["Metadata"]["B1"].value == "Dubai Land Department"
    assert workbook["Metadata"]["B5"].value == 1
    assert workbook["Metadata"]["B6"].value == transaction_service.MAX_TRANSACTION_EXPORT_ROWS
    assert workbook["Sales Transactions"]["A1"].value == "Dubai Land Department"
    assert workbook["Sales Transactions"]["A6"].value == "Transaction ID"
    assert workbook["Sales Transactions"]["A7"].value == "T-1"


async def test_sales_csv_export_returns_raw_data_without_workbook_metadata(monkeypatch):
    async def fake_query(sql, params):
        return [SALES_ROW]

    monkeypatch.setattr(transaction_service, "query", fake_query)

    stream = await transaction_service.export_sales_transactions(
        area="Dubai Marina",
        file_format="csv",
    )

    content = stream.getvalue().decode("utf-8-sig")
    assert content.startswith("Transaction ID,Date,Usage")
    assert "Dubai Land Department" not in content
    assert "T-1,2026-05-01,Residential" in content


async def test_sales_exports_escape_spreadsheet_formula_strings(monkeypatch):
    dangerous_row = {
        **SALES_ROW,
        "project_name_en": '=HYPERLINK("https://example.com")',
        "building_name_en": "+SUM(1,1)",
        "nearest_landmark_en": "-2+3",
        "nearest_metro_en": "@malicious",
        "nearest_mall_en": "\t=cmd",
    }

    async def fake_query(sql, params):
        return [dangerous_row]

    monkeypatch.setattr(transaction_service, "query", fake_query)

    csv_stream = await transaction_service.export_sales_transactions(file_format="csv")
    csv_content = csv_stream.getvalue().decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_content
    assert "'+SUM(1,1)" in csv_content
    assert "'-2+3" in csv_content
    assert "'@malicious" in csv_content
    assert "'\t=cmd" in csv_content

    excel_stream = await transaction_service.export_sales_transactions(file_format="xlsx")
    workbook = load_workbook(excel_stream)
    worksheet = workbook["Sales Transactions"]
    assert worksheet["N7"].value.startswith("'=HYPERLINK")
    assert worksheet["O7"].value == "'+SUM(1,1)"
    assert worksheet["P7"].value == "'-2+3"
    assert worksheet["Q7"].value == "'@malicious"
    assert worksheet["R7"].value == "'\t=cmd"


async def test_rental_excel_export_has_dld_metadata_and_50000_row_limit(monkeypatch):
    seen = {}

    async def fake_query(sql, params):
        seen["sql"] = sql
        seen["params"] = params
        return [RENTAL_ROW]

    monkeypatch.setattr(transaction_service, "query", fake_query)

    stream = await transaction_service.export_rental_contracts(area="Dubai Marina")

    assert "LIMIT {export_limit:Int32}" in seen["sql"]
    assert seen["params"]["export_limit"] == transaction_service.MAX_TRANSACTION_EXPORT_ROWS

    workbook = load_workbook(stream)
    assert workbook.sheetnames[:2] == ["Metadata", "Rental Contracts"]
    assert workbook["Metadata"]["B1"].value == "Dubai Land Department"
    assert workbook["Rental Contracts"]["A6"].value == "Contract ID"
    assert workbook["Rental Contracts"]["A7"].value == "R-1"
